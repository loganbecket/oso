from pathlib import Path

from oso import convert, db, handwriting, index, notes
from oso.config import Config, Course


def make_cfg(tmp_path: Path) -> Config:
    vault = tmp_path / "vault"
    (vault / "Courses" / "Calculus I" / "Lectures").mkdir(parents=True)
    (vault / "Inbox" / "Handwriting").mkdir(parents=True)
    return Config(vault=vault, timezone="America/New_York", courses=[Course("MATH-101-001", "Calculus I", "Calculus I")])


def test_front_matter_roundtrip():
    text = notes.with_front_matter({"type": "notes", "course": "MATH-101-001", "skip": None}, "# Hi\n\nbody")
    fm, body = notes.read_front_matter(text)
    assert fm == {"type": "notes", "course": "MATH-101-001"}
    assert body.startswith("# Hi")


def test_convert_docx_pptx_xlsx_pdf(tmp_path: Path):
    import docx
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfWriter

    cfg = make_cfg(tmp_path)
    lect = cfg.vault / "Courses" / "Calculus I" / "Lectures"

    d = docx.Document()
    d.add_heading("Limits", level=1)
    d.add_paragraph("A limit describes the value a function approaches.")
    d.save(str(lect / "Week 1 notes.docx"))

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Derivatives"
    slide.placeholders[1].text = "Rate of change"
    prs.save(str(lect / "Lecture 2.pptx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    ws.append(["Item", "Weight"])
    ws.append(["Homework", 30])
    wb.save(str(lect / "weights.xlsx"))

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with (lect / "blank.pdf").open("wb") as f:
        w.write(f)

    written = convert.convert_vault(cfg)
    names = sorted(p.name for p in written)
    assert names == ["Lecture 2.pptx.md", "Week 1 notes.docx.md", "blank.pdf.md", "weights.xlsx.md"]

    md = (lect / "Week 1 notes.docx.md").read_text()
    fm, body = notes.read_front_matter(md)
    assert fm["course"] == "MATH-101-001" and fm["type"] == "lecture"
    assert "## Limits" in body and "approaches" in body
    assert "Rate of change" in (lect / "Lecture 2.pptx.md").read_text()
    assert "| Homework | 30 |" in (lect / "weights.xlsx.md").read_text()
    assert "No text layer" in (lect / "blank.pdf.md").read_text()

    # Second run converts nothing: originals are not newer than their Markdown.
    assert convert.convert_vault(cfg) == []


def test_index_and_search(tmp_path: Path):
    cfg = make_cfg(tmp_path)
    note = cfg.vault / "Courses" / "Calculus I" / "Lectures" / "Lecture 3.md"
    note.write_text("---\ntype: lecture\n---\n# Chain rule\n\nThe chain rule differentiates composite functions.\n\n## Example\n\nd/dx sin(x^2) = 2x cos(x^2)\n")
    (cfg.vault / "Today.md").write_text("# Today\n\nchain rule should not be indexed here")
    with db.connect(tmp_path / "t.sqlite") as conn:
        counts = index.rebuild(conn, cfg)
        assert counts["indexed"] == 1
        hits = index.search(conn, "chain rule composite", course="MATH-101-001")
        assert hits and hits[0]["path"] == "Courses/Calculus I/Lectures/Lecture 3.md"
        assert hits[0]["type"] == "lecture"
        assert index.search(conn, "chain rule", course="PHYS-110") == []
        assert index.rebuild(conn, cfg)["indexed"] == 0
        note.unlink()
        assert index.rebuild(conn, cfg)["removed"] == 1
        assert index.search(conn, "chain rule") == []


def test_handwriting_queue_and_mark(tmp_path: Path):
    from pypdf import PdfWriter

    cfg = make_cfg(tmp_path)
    w = PdfWriter()
    w.add_blank_page(width=300, height=400)
    w.add_blank_page(width=300, height=400)
    with (cfg.vault / "Inbox" / "Handwriting" / "Physics week 3.pdf").open("wb") as f:
        w.write(f)
    with db.connect(tmp_path / "t.sqlite") as conn:
        assert handwriting.queue_new(conn, cfg) == 2
        assert handwriting.queue_new(conn, cfg) == 0
        pend = handwriting.pending(conn)
        assert [p["page"] for p in pend] == [1, 2]
        assert (cfg.vault / pend[0]["path"]).exists()
        handwriting.mark(conn, pend[0]["path"], "Courses/Physics/Notes/x.md", 0.4)
        handwriting.mark(conn, pend[1]["path"], "Courses/Physics/Notes/x.md", 0.95)
        assert handwriting.pending(conn) == []
        low = handwriting.low_confidence(conn)
        assert len(low) == 1 and low[0]["page"] == 1
