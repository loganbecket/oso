"""Convert Word, PowerPoint, Excel, and PDF files to Markdown beside the original.

The Markdown file has the same name with `.md` appended, so `Lecture 3.pptx` gets `Lecture 3.pptx.md`.
A file is reconverted only when the original is newer than its Markdown.
"""

from __future__ import annotations

import logging
from pathlib import Path

from . import notes
from .config import Config

log = logging.getLogger("oso.convert")

CONVERTIBLE = {".docx", ".pptx", ".xlsx", ".pdf"}


def convert_vault(cfg: Config) -> list[Path]:
    """Walk Courses/ and Inbox/ and convert anything new. Returns the Markdown files written."""
    written: list[Path] = []
    for root in (cfg.vault / "Courses", cfg.vault / "Inbox"):
        if not root.exists():
            continue
        for src in root.rglob("*"):
            if src.suffix.lower() not in CONVERTIBLE or not src.is_file():
                continue
            if "Handwriting" in src.parts:
                continue  # handwritten pages are transcribed, not converted
            out = src.with_name(src.name + ".md")
            if out.exists() and out.stat().st_mtime >= src.stat().st_mtime:
                continue
            try:
                body = to_markdown(src)
            except Exception as e:  # noqa: BLE001
                log.warning("could not convert %s: %s", src.name, type(e).__name__)
                continue
            course = _course_from_path(cfg, src)
            fm = {
                "type": notes.guess_type(src.stem),
                "course": course,
                "source": src.name,
                "converted": notes.stamp(),
            }
            out.write_text(notes.with_front_matter(fm, f"# {src.stem}\n\n{body}"), encoding="utf-8")
            written.append(out)
    return written


def to_markdown(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".xlsx":
        return _xlsx(path)
    if ext == ".pdf":
        return _pdf(path)
    raise ValueError(f"unsupported file type {ext}")


def _docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    out: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            out.append(f"{'#' * min(int(level) + 1, 6)} {text}")
        elif style.startswith("list"):
            out.append(f"- {text}")
        else:
            out.append(text)
        out.append("")
    for t in d.tables:
        out.append(_table([[c.text.strip() for c in row.cells] for row in t.rows]))
        out.append("")
    return "\n".join(out)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## Slide {n}{': ' + title if title else ''}")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        out.append(f"{'  ' * para.level}- {text}")
            if getattr(shape, "has_table", False) and shape.has_table:
                out.append(_table([[c.text.strip() for c in row.cells] for row in shape.table.rows]))
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                out.append(f"\n> Speaker notes: {notes_text}")
        out.append("")
    return "\n".join(out)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True, max_row=200)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        out.append(f"## Sheet: {ws.title}")
        out.append(_table(rows))
        out.append("")
    return "\n".join(out)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[str] = []
    for n, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append(f"## Page {n}\n\n{text}\n")
    if not out:
        out.append("(No text layer found. If this is scanned or handwritten, move it to Inbox/Handwriting to have it transcribed.)")
    return "\n".join(out)


def _table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    esc = lambda c: c.replace("|", "\\|").replace("\n", " ")  # noqa: E731
    head = "| " + " | ".join(esc(c) for c in rows[0]) + " |"
    sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
    body = ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join([head, sep, *body])


def _course_from_path(cfg: Config, path: Path) -> str | None:
    try:
        rel = path.relative_to(cfg.vault / "Courses")
    except ValueError:
        return None
    folder = rel.parts[0] if rel.parts else None
    for c in cfg.courses:
        if c.folder == folder:
            return c.code
    return None
