from pathlib import Path

from oso import convert, db, handwriting, notes
from oso.config import Config, Course


def inked_pdf(path, pages=1):
    """A PDF whose pages carry ink, so the blank-page filter keeps them."""
    from PIL import Image, ImageDraw

    imgs = []
    for _ in range(pages):
        im = Image.new("L", (600, 800), 255)
        ImageDraw.Draw(im).rectangle((80, 80, 520, 200), fill=0)
        imgs.append(im)
    imgs[0].save(path, save_all=True, append_images=imgs[1:])


def make_cfg(tmp_path: Path) -> Config:
    vault = tmp_path / "vault"
    (vault / "Courses" / "Calculus I" / "Lectures").mkdir(parents=True)
    (vault / "Courses" / "Calculus I" / "Handwriting").mkdir(parents=True)
    return Config(vault=vault, timezone="America/New_York", courses=[Course("MATH-101-001", "Calculus I", "Calculus I")])


def test_front_matter_roundtrip():
    text = notes.with_front_matter({"type": "notes", "course": "MATH-101-001", "skip": None}, "# Hi\n\nbody")
    fm, body = notes.read_front_matter(text)
    assert fm == {"type": "notes", "course": "MATH-101-001"}
    assert body.startswith("# Hi")


def _odt(path):
    import zipfile

    content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">
 <office:body><office:text>
  <text:h text:outline-level="1">Kinematics</text:h>
  <text:p>Velocity is the derivative of position.</text:p>
 </office:text></office:body></office:document-content>"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text", compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", content)
        z.writestr("META-INF/manifest.xml", manifest)


def test_convert_office_pdf_libreoffice_and_google(tmp_path: Path):
    import docx
    from openpyxl import Workbook
    from pptx import Presentation

    cfg = make_cfg(tmp_path)
    lect = cfg.vault / "Courses" / "Calculus I" / "Lectures"

    d = docx.Document()
    d.add_heading("Limits", level=1)
    d.add_paragraph("A limit describes the value a function approaches.")
    d.save(str(lect / "Week 1 notes.docx"))

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Derivatives"
    slide.placeholders[1].text = "Rate of change"
    prs.save(str(lect / "Lecture 2.pptx"))

    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Weight"])
    ws.append(["Homework", 30])
    wb.save(str(lect / "weights.xlsx"))

    inked_pdf(lect / "scan.pdf")
    _odt(lect / "Kinematics.odt")
    (lect / "Shared notes.gdoc").write_text('{"doc_id": "abc123", "email": "x@y"}')

    written = convert.convert_vault(cfg)
    names = sorted(p.name for p in written)
    assert names == ["Kinematics.odt.md", "Lecture 2.pptx.md", "Shared notes.gdoc.md", "Week 1 notes.docx.md", "scan.pdf.md", "weights.xlsx.md"]

    fm, body = notes.read_front_matter((lect / "Week 1 notes.docx.md").read_text())
    assert fm["course"] == "MATH-101-001" and fm["type"] == "lecture"
    assert "Limits" in body and "approaches" in body
    assert "Rate of change" in (lect / "Lecture 2.pptx.md").read_text()
    assert "Homework" in (lect / "weights.xlsx.md").read_text()
    assert "No text layer" in (lect / "scan.pdf.md").read_text()
    assert "Velocity is the derivative" in (lect / "Kinematics.odt.md").read_text()
    g = (lect / "Shared notes.gdoc.md").read_text()
    assert "type: google-file" in g and "https://docs.google.com/document/d/abc123" in g

    assert convert.convert_vault(cfg) == []


def test_opendocument_without_libreoffice(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(convert, "libreoffice", lambda: None)
    _odt(tmp_path / "k.odt")
    text = convert.to_markdown(tmp_path / "k.odt")
    assert "## Kinematics" in text and "Velocity is the derivative" in text


def test_split_sections():
    secs = notes.split_sections("intro\n# A\n\none\n## B\n\ntwo")
    assert secs == [("", "intro"), ("A", "one"), ("B", "two")]


def test_handwriting_queue_and_mark(tmp_path: Path):
    cfg = make_cfg(tmp_path)
    inked_pdf(cfg.vault / "Courses" / "Calculus I" / "Handwriting" / "Week 3.pdf", pages=2)
    stray = cfg.vault / "Clippings"
    stray.mkdir()
    inked_pdf(stray / "No course.pdf")  # scans outside a course folder are ignored
    with db.connect(tmp_path / "t.sqlite") as conn:
        assert handwriting.queue_new(conn, cfg) == 2
        assert handwriting.queue_new(conn, cfg) == 0
        pend = handwriting.pending(conn)
        assert [p["page"] for p in pend] == [1, 2]
        assert (cfg.vault / pend[0]["path"]).exists()
        handwriting.mark(conn, pend[0]["path"], "Courses/Physics/Notes/x.md", 0.4)
        handwriting.mark(conn, pend[1]["path"], "Courses/Physics/Notes/x.md", 0.95)
        assert handwriting.pending(conn) == []
        low = handwriting.low_confidence(conn)
        assert len(low) == 1 and low[0]["page"] == 1


def test_handwriting_without_extension_is_recognized(tmp_path: Path):
    cfg = make_cfg(tmp_path)
    hw = cfg.vault / "Courses" / "Calculus I" / "Handwriting"
    inked_pdf(hw / "scan.pdf")
    (hw / "scan.pdf").rename(hw / "3f9a1c0e7b2d4a8f")
    (hw / "notes.txt").write_text("not a scan")
    (hw / "a1b2c3").write_text("not a scan either")
    with db.connect(tmp_path / "t.sqlite") as conn:
        assert handwriting.queue_new(conn, cfg) == 1
